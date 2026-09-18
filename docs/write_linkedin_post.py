#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(13)
font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("SwarmAI — A LinkedIn Post")
run.bold = True
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

doc.add_paragraph("")  # spacer

doc.add_paragraph("SwarmAI — a native coding agent app for Mac, built entirely in Swift & SwiftUI.").italic = True

post_text = """I built SwarmAI because I was tired of bouncing between terminal windows, git repos, and chat UIs just to get things done with my coding agents.

We all have these amazing tools now — Claude Code, Codex, Cursor, Copilot, DeepSeek — but there wasn't really one place to hold them all together on a Mac. So I built one.

Here's what makes it different:

It doesn't replace your agents. It sits on top of the ones you already use, the ones you're already logged into with your own subscriptions. Codex, Claude, Cursor, Copilot, DeepSeek, Meta, Grok — all in one app, no new API keys to buy, no lock-in.

The UI is built in native SwiftUI with Liquid Glass, because honestly, developer tools don't have to look like they're from 2012. It feels right at home on a Mac.

The workflow is what I cared about most:
• Projects and conversations in a glass sidebar — pin the ones you revisit, let the rest settle to the bottom.
• Every response streams in with reasoning, tool calls, and a per-turn diff you can review and revert. No more wondering what your agent actually changed.
• Embedded terminal per thread, so you can inspect or run things without leaving the conversation.
• Commit, push, and open PRs — generated messages, no copy-pasting.
• Hydra mode — one agent leads a team of helpers, each on the model and effort you pick, even across providers.
• Plan mode, four permission levels, model and reasoning selection, 26 themes including Catppuccin, Dracula, Claude, and Codex.

Built for Apple Silicon on macOS 26 and later. Remote access, cloud sync, telemetry, and a web client are left out on purpose. Your code stays on your machine.

If this sounds like something you'd use, the repo and a signed Apple Silicon build are both live.

Would love to hear from anyone who's been thinking the same thing.

#BuildInPublic #SwiftUI #DeveloperTools #MacApp #OpenSource"""

for para_text in post_text.split("\n\n"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(para_text)
    run.font.size = Pt(13)

# Hashtags row
doc.add_paragraph("")
htag = doc.add_paragraph()
htag.paragraph_format.space_after = Pt(6)
run = htag.add_run("#BuildInPublic  #SwiftUI  #DeveloperTools  #MacApp  #OpenSource")
run.font.size = Pt(13)
run.bold = True

doc.save("/Users/soumyachakraborty/Documents/droppy-code/docs/linkedin_swarmai_post.docx")
print("Saved.")
